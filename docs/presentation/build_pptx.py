#!/usr/bin/env python3
"""Build the ADAD Test Automation Framework deck in MNC visual language."""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ASSETS = Path(__file__).resolve().parent / "assets"
OUT = Path(__file__).resolve().parent.parent / "ADAD_Test_Automation_Framework.pptx"

# MNC brand (sampled from the guideline deck)
BLUE = RGBColor(37, 105, 237)
BLUE_SOFT = RGBColor(232, 240, 255)
BLUE_MID = RGBColor(90, 140, 242)
NAVY = RGBColor(22, 36, 71)
BODY = RGBColor(45, 55, 72)
MUTED = RGBColor(100, 116, 139)
WHITE = RGBColor(255, 255, 255)
CODE_BG = RGBColor(22, 32, 58)
CODE_FG = RGBColor(226, 232, 240)
CODE_KEY = RGBColor(147, 197, 253)
CODE_STR = RGBColor(167, 243, 208)
CODE_CMT = RGBColor(148, 163, 184)
CODE_FN = RGBColor(253, 224, 71)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FOOTER_H = Inches(0.88)
MARGIN = Inches(0.72)
FONT = "Calibri"
FONT_CODE = "Consolas"

KEYWORDS = {
    "export",
    "class",
    "extends",
    "constructor",
    "super",
    "async",
    "await",
    "new",
    "from",
    "import",
    "const",
    "return",
    "public",
    "void",
    "string",
    "static",
    "readonly",
    "test",
}


def rgb_of(color: RGBColor) -> str:
    return f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"


def set_run_font(run, name: str, size: Pt, bold: bool = False, color: RGBColor | None = None, italic: bool = False):
    run.font.name = name
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:cs", "a:ea"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", name)


def set_slide_bg(slide, color: RGBColor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, fill: RGBColor, line: RGBColor | None = None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    shape.shadow.inherit = False
    return shape


def add_round(slide, left, top, width, height, fill: RGBColor, line: RGBColor | None = None, adj: int = 4500):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1.25)
    shape.shadow.inherit = False
    prst = shape._element.spPr.find(qn("a:prstGeom"))
    av_lst = prst.find(qn("a:avLst"))
    if av_lst is None:
        av_lst = etree.SubElement(prst, qn("a:avLst"))
    else:
        for child in list(av_lst):
            av_lst.remove(child)
    gd = etree.SubElement(av_lst, qn("a:gd"))
    gd.set("name", "adj")
    gd.set("fmla", f"val {adj}")
    return shape


def textbox(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    return tf


def write_p(p, text, size, color, bold=False, align=PP_ALIGN.LEFT, space_after=6, space_before=0, font=FONT):
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.clear()
    run = p.add_run()
    run.text = text
    set_run_font(run, font, size, bold=bold, color=color)
    return p


def add_label(slide, left, top, width, height, text, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tf = textbox(slide, left, top, width, height, anchor=anchor)
    write_p(tf.paragraphs[0], text, size, color, bold=bold, align=align, space_after=0)
    return tf


def add_footer(slide, page: int, total: int):
    add_rect(slide, 0, SLIDE_H - FOOTER_H, SLIDE_W, FOOTER_H, BLUE)
    add_label(
        slide,
        0,
        SLIDE_H - FOOTER_H,
        SLIDE_W,
        FOOTER_H,
        str(page),
        Pt(14),
        WHITE,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    logo = ASSETS / "mc_footer.png"
    if logo.exists():
        slide.shapes.add_picture(str(logo), Inches(12.05), SLIDE_H - Inches(0.70), height=Inches(0.50))


def add_confidential(slide):
    add_label(
        slide,
        Inches(10.6),
        Inches(0.18),
        Inches(2.3),
        Inches(0.32),
        "CONFIDENTIAL",
        Pt(12),
        BLUE,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )


def add_title(slide, text: str):
    add_label(slide, MARGIN, Inches(0.28), Inches(11.5), Inches(0.72), text, Pt(32), BLUE, bold=True)


def add_bullets(slide, left, top, width, height, items: list[str], size=Pt(18)):
    tf = textbox(slide, left, top, width, height)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(8)
        p.space_after = Pt(8)
        dash = p.add_run()
        dash.text = "–  "
        set_run_font(dash, FONT, size, bold=True, color=BLUE)
        body = p.add_run()
        body.text = item
        set_run_font(body, FONT, size, color=BODY)


def highlight_code(p, line: str, size=Pt(13)):
    p.space_before = Pt(1)
    p.space_after = Pt(1)
    p.line_spacing = 1.12
    stripped = line.rstrip("\n")
    comment = ""
    hash_at = stripped.find("#")
    if stripped.lstrip().startswith("#") or stripped.lstrip().startswith("//"):
        run = p.add_run()
        run.text = stripped
        set_run_font(run, FONT_CODE, size, color=CODE_CMT, italic=True)
        return
    if hash_at != -1:
        comment = stripped[hash_at:]
        stripped = stripped[:hash_at]
    slash = stripped.find("//")
    if slash != -1:
        comment = stripped[slash:] + comment
        stripped = stripped[:slash]
    token = ""
    in_str = False
    quote = ""
    i = 0
    while i < len(stripped):
        ch = stripped[i]
        if in_str:
            token += ch
            if ch == quote and stripped[i - 1] != "\\":
                run = p.add_run()
                run.text = token
                set_run_font(run, FONT_CODE, size, color=CODE_STR)
                token = ""
                in_str = False
            i += 1
            continue
        if ch in ('"', "'", "`"):
            if token:
                emit_code_token(p, token, size)
                token = ""
            in_str = True
            quote = ch
            token = ch
            i += 1
            continue
        if ch.isspace() or ch in "(){},.=:;<>[]":
            if token:
                emit_code_token(p, token, size)
                token = ""
            run = p.add_run()
            run.text = ch
            set_run_font(run, FONT_CODE, size, color=CODE_FG)
            i += 1
            continue
        token += ch
        i += 1
    if token:
        emit_code_token(p, token, size)
    if comment:
        run = p.add_run()
        run.text = comment
        set_run_font(run, FONT_CODE, size, color=CODE_CMT, italic=True)


def emit_code_token(p, token: str, size):
    run = p.add_run()
    run.text = token
    if token in KEYWORDS:
        set_run_font(run, FONT_CODE, size, color=CODE_KEY, bold=True)
    elif token[:1].isupper() and token.isalnum():
        set_run_font(run, FONT_CODE, size, color=CODE_FN)
    else:
        set_run_font(run, FONT_CODE, size, color=CODE_FG)


def add_code(slide, left, top, width, height, code: str, size=Pt(13)):
    add_round(slide, left, top, width, height, CODE_BG, adj=3500)
    tf = textbox(slide, left + Inches(0.28), top + Inches(0.18), width - Inches(0.48), height - Inches(0.32))
    lines = code.strip("\n").split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        highlight_code(p, line if line else " ", size)


def card(slide, left, top, width, height, title: str, body: str):
    add_round(slide, left, top, width, height, BLUE_SOFT, adj=4000)
    add_rect(slide, left, top, Inches(0.10), height, BLUE)
    add_label(slide, left + Inches(0.28), top + Inches(0.18), width - Inches(0.42), Inches(0.40), title, Pt(16), BLUE, bold=True)
    tf = textbox(slide, left + Inches(0.28), top + Inches(0.58), width - Inches(0.42), height - Inches(0.72))
    write_p(tf.paragraphs[0], body, Pt(14), BODY, space_after=0)


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    total = 12

    # 1 — Title
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, BLUE)
    s.shapes.add_picture(str(ASSETS / "logo_white.png"), Inches(0.70), Inches(0.42), width=Inches(1.95))
    s.shapes.add_picture(str(ASSETS / "planes.png"), Inches(6.55), Inches(1.35), width=Inches(6.55))
    add_label(s, MARGIN, Inches(2.35), Inches(7.6), Inches(1.15), "Test Automation Framework", Pt(36), WHITE, bold=True)
    add_label(s, MARGIN, Inches(3.45), Inches(6.4), Inches(0.50), "ADAD  ·  Playwright  ·  TypeScript", Pt(22), WHITE)
    add_label(s, MARGIN, Inches(4.15), Inches(6.4), Inches(0.40), "Quality Assurance", Pt(16), WHITE)
    add_label(s, MARGIN, Inches(6.35), Inches(4.0), Inches(0.40), "14/09/2026", Pt(16), WHITE)

    # 2 — Why
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, WHITE)
    add_confidential(s)
    add_title(s, "Why this framework?")
    add_label(
        s,
        MARGIN,
        Inches(1.10),
        Inches(11.8),
        Inches(0.55),
        "A scalable, maintainable suite for ADAD: one page object per route, reusable TestPoints, and Playwright specs that compose them.",
        Pt(16),
        BODY,
    )
    card(
        s,
        MARGIN,
        Inches(1.80),
        Inches(3.85),
        Inches(2.35),
        "Page Object Model",
        "Each ADAD route has one page class that owns locators and the TestPoints that use them. UI changes stay in one file.",
    )
    card(
        s,
        Inches(4.75),
        Inches(1.80),
        Inches(3.85),
        Inches(2.35),
        "TestPoints",
        "Reusable steps with an id, description and ticket. Specs assemble TestPoints instead of duplicating clicks and assertions.",
    )
    card(
        s,
        Inches(8.78),
        Inches(1.80),
        Inches(3.85),
        Inches(2.35),
        "Playwright",
        "Typed tests, worker-scoped login, HTML/JSON reports, traces and screenshots on failure, ready for CI.",
    )
    add_label(s, MARGIN, Inches(4.35), Inches(11.8), Inches(0.35), "Also in the suite", Pt(16), BLUE, bold=True)
    add_bullets(
        s,
        MARGIN,
        Inches(4.70),
        Inches(11.8),
        Inches(1.55),
        [
            "YAML + .env configuration, Keycloak login, Postgres access via SSM tunnel, SFTP file injection",
            "47 page objects covering dashboard, matching, audit track, alerts, accounting and settings",
            "Clear split: src/ owns automation building blocks, tests/ is the only place tests run",
        ],
        size=Pt(16),
    )
    add_footer(s, 2, total)

    # 3 — Architecture as stacked layers
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, WHITE)
    add_confidential(s)
    add_title(s, "Architecture")
    add_label(
        s,
        MARGIN,
        Inches(1.02),
        Inches(8.5),
        Inches(0.32),
        "Each layer only talks to the one below.",
        Pt(14),
        MUTED,
    )
    stack = [
        (RGBColor(85, 145, 244), "Playwright specs", "smoke · filters · navigation"),
        (RGBColor(52, 118, 238), "runTestPoints()", "Runs TestPoints in order on the logged-in page"),
        (BLUE, "TestPoints", "Reusable action + id + ticket"),
        (RGBColor(28, 58, 140), "Page objects", "Locators and UI methods — one class per ADAD route"),
        (NAVY, "Foundation", "fixtures · navigation · configuration · SFTP / DB · Playwright"),
    ]
    bar_h = Inches(0.78)
    gap = Inches(0.10)
    y = Inches(1.40)
    slide_center = SLIDE_W.inches / 2
    for i, (fill, title, sub) in enumerate(stack):
        width = Inches(8.15 + i * 0.95)
        left = Inches(slide_center - width.inches / 2)
        add_round(s, left, y, width, bar_h, fill, adj=2800)
        add_label(
            s,
            left + Inches(0.28),
            y,
            Inches(3.35),
            bar_h,
            title,
            Pt(16),
            WHITE,
            bold=True,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        add_label(
            s,
            left + Inches(3.55),
            y,
            width - Inches(3.80),
            bar_h,
            sub,
            Pt(13),
            WHITE,
            align=PP_ALIGN.RIGHT,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        y += bar_h + gap
    add_footer(s, 3, total)

    # 4 — Folder structure
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, WHITE)
    add_confidential(s)
    add_title(s, "Folder structure")
    add_code(
        s,
        MARGIN,
        Inches(1.20),
        Inches(6.15),
        Inches(4.95),
        """
src/
  pages/            # one .page.ts per ADAD route
  testpoint/        # TestPoint base class
  navigation/       # URL navigation TestPoints
  configuration/    # yaml, env, login, DB, SFTP
  fixtures/         # worker-scoped logged-in page
tests/
  runTestPoints.ts  # shared TestPoint runner
  smoke/            # ingest + matching journeys
  filters/          # audit-track filter matrices
  navigation/       # route-by-link checks
  testpoints/       # registry + live inspect
  framework/        # TestPoint contract tests
docs/test_files/    # CSV / JSONL fixtures
scripts/            # page generation, DB tunnel
""".strip(),
        size=Pt(13),
    )
    add_label(s, Inches(7.20), Inches(1.25), Inches(5.4), Inches(0.40), "How to read it", Pt(18), BLUE, bold=True)
    add_bullets(
        s,
        Inches(7.20),
        Inches(1.70),
        Inches(5.4),
        Inches(4.3),
        [
            "pages/ is the Page Object Model: locators live next to the TestPoints that use them.",
            "tests/ is the only executable layer. Specs import TestPoints and call runTestPoints().",
            "navigation/ covers every mapped ADAD URL as a reusable step.",
            "configuration/ keeps secrets in .env and shared paths in config.yaml.",
            "The stack is Playwright and TypeScript end to end.",
        ],
        size=Pt(15),
    )
    add_footer(s, 4, total)

    # 5 — TestPoints
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, WHITE)
    add_confidential(s)
    add_title(s, "TestPoints layer")
    add_label(
        s,
        MARGIN,
        Inches(1.10),
        Inches(11.8),
        Inches(0.70),
        "A TestPoint is one reusable action: open a page, click Search, inject a file, assert a match. Specs are lists of TestPoints, not duplicated locators and clicks.",
        Pt(16),
        BODY,
    )
    benefits = [
        ("Reusable", "The same step is shared by smoke, filters and navigation without rewriting locators."),
        ("Traceable", "Each step carries an id, a description and a ticket id for reporting and Jira."),
        ("Maintainable", "When the UI moves, you fix the page object once. Every spec that uses that TestPoint follows."),
        ("Composable", "Journeys are assembled with runTestPoints([new StepA(), new StepB(), ...])."),
    ]
    positions = [
        (MARGIN, Inches(1.95)),
        (Inches(7.05), Inches(1.95)),
        (MARGIN, Inches(3.85)),
        (Inches(7.05), Inches(3.85)),
    ]
    for (left, top), (title, body) in zip(positions, benefits):
        card(s, left, top, Inches(5.55), Inches(1.65), title, body)
    add_footer(s, 5, total)

    # 6 — TestPoint example
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, WHITE)
    add_confidential(s)
    add_title(s, "Example — TestPoint")
    add_code(
        s,
        MARGIN,
        Inches(1.15),
        Inches(12.0),
        Inches(4.70),
        """
export class OpenManualMatchingPage extends TestPoint {
  constructor() {
    super(
      "TP-OPEN-MANUAL_MATCHING",
      "Open /manual-matching",
      async (page) => {
        const pom = new ManualMatchingPage(page);
        await pom.goto();
        await pom.expectOnPage();
      },
    );
  }
}
""".strip(),
        size=Pt(16),
    )
    add_footer(s, 6, total)

    # 7 — Spec example
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, WHITE)
    add_confidential(s)
    add_title(s, "Example — composing a spec")
    add_code(
        s,
        MARGIN,
        Inches(1.15),
        Inches(12.0),
        Inches(4.70),
        """
test("verify the manual matching", async ({ loggedInPage }) => {
  await runTestPoints(loggedInPage, [
    new InjectFile(salePath, adadPaths.sales, saleFileName),
    new TestPointNavigateToManualMatchingByLink(),
    new ManualMatchingSaleUIFilterPNROrderReference("R4Y7W4"),
    new ManualMatchingUIConfirmMatching(),
    new ManualMatchingUIMatchSaleAndSettlement(),
    new SaleUICheckIfMatchingDoneCorrectly("1"),
  ]);
});
""".strip(),
        size=Pt(15),
    )
    add_footer(s, 7, total)

    # 8 — POM
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, WHITE)
    add_confidential(s)
    add_title(s, "Page Object Model")
    add_label(
        s,
        MARGIN,
        Inches(1.10),
        Inches(11.8),
        Inches(0.55),
        "One file per route under src/pages/. The class holds locators; exported TestPoint classes perform the actions.",
        Pt(16),
        BODY,
    )
    add_code(
        s,
        MARGIN,
        Inches(1.75),
        Inches(7.35),
        Inches(4.35),
        """
export class ManualMatchingPage extends BasePage {
  static readonly path = "/manual-matching";

  constructor(page: Page) {
    super(page, ManualMatchingPage.path);
    this.search = this.findNamedButton("Search");
    this.confirm = this.page.locator(
      "button.ConfirmButton"
    );
  }
}
""".strip(),
        size=Pt(14),
    )
    add_label(s, Inches(8.30), Inches(1.75), Inches(4.3), Inches(0.40), "Benefits", Pt(18), BLUE, bold=True)
    add_bullets(
        s,
        Inches(8.30),
        Inches(2.20),
        Inches(4.3),
        Inches(3.8),
        [
            "Reusable locators via BasePage helpers (tabs, dropdowns, date pickers, export).",
            "UI changes stay in the page class, not in every spec.",
            "TestPoints on the same file keep action and selector together.",
            "Inventory lists every route so registry tests can inspect coverage.",
        ],
        size=Pt(15),
    )
    add_footer(s, 8, total)

    # 9 — Execution & reporting
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, WHITE)
    add_confidential(s)
    add_title(s, "Execution and reporting")
    exec_cards = [
        ("VS Code", "Playwright Test panel or Testing view. Run a spec, a TestPoint health check, or the full suite headed / UI mode."),
        ("Command line", "npx playwright test tests/smoke\nnpm run test:filters\nnpm run test:testpoints"),
        ("CI / CD", "playwright.config.ts already sets CI retries, HTML + JSON reporters, traces, screenshots and video on failure."),
    ]
    x = MARGIN
    for title, body in exec_cards:
        add_round(s, x, Inches(1.20), Inches(3.85), Inches(2.35), BLUE_SOFT, adj=4000)
        add_rect(s, x, Inches(1.20), Inches(3.85), Inches(0.12), BLUE)
        add_label(s, x + Inches(0.22), Inches(1.45), Inches(3.40), Inches(0.40), title, Pt(18), BLUE, bold=True)
        tf = textbox(s, x + Inches(0.22), Inches(1.90), Inches(3.40), Inches(1.45))
        for i, line in enumerate(body.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            write_p(p, line, Pt(13), BODY, space_after=4, font=FONT_CODE if line.startswith("npx") or line.startswith("npm") else FONT)
        x += Inches(4.05)

    reports = [
        ("HTML report", "playwright-report/ — timeline, traces, screenshots"),
        ("JSON results", "reports/playwright-results.json"),
        ("Per-test logs", "logs/<date>/… with TestPoint start / fail lines"),
        ("Failure artefacts", "trace, screenshot and video retained on fail"),
    ]
    y = Inches(3.80)
    for title, body in reports:
        add_round(s, MARGIN, y, Inches(11.90), Inches(0.58), BLUE_SOFT, adj=2500)
        add_label(s, MARGIN + Inches(0.25), y, Inches(2.6), Inches(0.58), title, Pt(15), BLUE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        add_label(s, Inches(3.70), y, Inches(8.6), Inches(0.58), body, Pt(15), BODY, anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(0.66)
    add_footer(s, 9, total)

    # 10 — Next challenges: data + versions
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, WHITE)
    add_confidential(s)
    add_title(s, "Next challenges")
    add_label(
        s,
        MARGIN,
        Inches(1.08),
        Inches(11.8),
        Inches(0.48),
        "The framework is in place. The interesting work is keeping tests true to ADAD: same data, same behaviour, every release.",
        Pt(16),
        BODY,
    )
    add_round(s, MARGIN, Inches(1.65), Inches(5.85), Inches(4.50), BLUE_SOFT, adj=4000)
    add_rect(s, MARGIN, Inches(1.65), Inches(0.10), Inches(4.50), BLUE)
    add_label(s, MARGIN + Inches(0.32), Inches(1.82), Inches(5.30), Inches(0.42), "Data reproduction", Pt(20), BLUE, bold=True)
    add_bullets(
        s,
        MARGIN + Inches(0.32),
        Inches(2.32),
        Inches(5.25),
        Inches(3.55),
        [
            "Rebuild a known molecule: sales, settlements, bank deposits, dates and currencies.",
            "Scenario packs already exist: correct, discrepancy, duplicate, late, sealed, many-to-many, EUR vs TND.",
            "Inject over SFTP, wait for ETL, then assert in the UI and in Postgres.",
            "Staging is shared — leftover files and matches pollute the next run. Isolation and cleanup are the next step.",
        ],
        size=Pt(14),
    )
    add_round(s, Inches(6.78), Inches(1.65), Inches(5.85), Inches(4.50), BLUE_SOFT, adj=4000)
    add_rect(s, Inches(6.78), Inches(1.65), Inches(0.10), Inches(4.50), BLUE)
    add_label(s, Inches(7.10), Inches(1.82), Inches(5.30), Inches(0.42), "Each new ADAD version", Pt(20), BLUE, bold=True)
    add_bullets(
        s,
        Inches(7.10),
        Inches(2.32),
        Inches(5.25),
        Inches(3.55),
        [
            "A release can move a locator, add a route, or change a matching / alert rule.",
            "47 page objects must track the product, not last month’s UI.",
            "Button catalog + gen-pages catch new controls; hand-maintained pages stay for journeys.",
            "The TestPoint registry should fail loudly when a version ships uncovered screens.",
        ],
        size=Pt(14),
    )
    add_footer(s, 10, total)

    # 11 — What makes it interesting
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, WHITE)
    add_confidential(s)
    add_title(s, "What makes it interesting")
    add_label(
        s,
        MARGIN,
        Inches(1.08),
        Inches(11.8),
        Inches(0.42),
        "This is not a static UI suite. Tests run against a live reconciliation product, with real timing, data and risk.",
        Pt(16),
        BODY,
    )
    interesting = [
        (
            "Shared QA, not a sandbox",
            "Keycloak, SFTP, SSM tunnel and one staging schema. Tests compete with other users and leftover data.",
        ),
        (
            "Async ingestion",
            "A file is not in ADAD when SFTP returns. ETL, File Dashboard and Trigger now decide when matching can start.",
        ),
        (
            "Filters from live Postgres",
            "Audit-track matrices are built from real rows, not hardcoded IDs — including empty and negative combinations.",
        ),
        (
            "One worker, one session",
            "Login is worker-scoped. Parallel runs would collide on files, matches and the date-range limit (max one month).",
        ),
        (
            "Domain depth",
            "FX, payment-order grouping, alerts, accounting periods. Next journeys: sanity-check alerts and ledger posting.",
        ),
        (
            "Coverage vs risk",
            "Inspect-only TestPoints for destructive buttons. CI still depends on the tunnel, staging and a clean molecule.",
        ),
    ]
    positions = [
        (MARGIN, Inches(1.62)),
        (Inches(4.75), Inches(1.62)),
        (Inches(8.78), Inches(1.62)),
        (MARGIN, Inches(3.85)),
        (Inches(4.75), Inches(3.85)),
        (Inches(8.78), Inches(3.85)),
    ]
    for (left, top), (title, body) in zip(positions, interesting):
        card(s, left, top, Inches(3.85), Inches(2.05), title, body)
    add_footer(s, 11, total)

    # 12 — Thank you
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, WHITE)
    s.shapes.add_picture(str(ASSETS / "handshake.png"), Inches(0.55), Inches(0.55), height=Inches(5.85))
    add_label(s, Inches(6.55), Inches(1.55), Inches(5.8), Inches(0.85), "Thank you.", Pt(40), BLUE, bold=True)
    add_label(s, Inches(6.55), Inches(2.65), Inches(5.8), Inches(0.90), "ADAD Test Automation\nPlaywright · TypeScript", Pt(22), NAVY)
    add_label(s, Inches(6.55), Inches(4.05), Inches(5.8), Inches(0.40), "mnc.aero", Pt(16), BLUE)
    add_footer(s, 12, total)

    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
